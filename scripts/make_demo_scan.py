"""Создаёт демонстрационный «скан» docs/samples/demo-scan.pdf без текстового слоя.

В документе есть оформление (заголовок по центру, жирный шрифт, табуляции,
абзацы по ширине с отступом) и намеренные опечатки — чтобы показать сохранение
вёрстки и подсветку ошибок. Запуск: .venv/bin/python scripts/make_demo_scan.py
"""

import asyncio
import io
import sys
import tempfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services import office  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "docs" / "samples" / "demo-scan.pdf"

# (текст, выравнивание, жирный, кегль)
CONTENT = [
    ("ООО «Пример»", "center", True, 14),
    ("ПРИКАЗ № 17", "center", True, 16),
    ("г. Москва\t15 октября 2026 г.", "tab", False, 12),
    ("О вводе в эксплуатацию сервиса распознования документов", "center", True, 12),
    ("В целях повышения эффективности электронного документооборота и сокращения времени "
     "обработки входящей корреспонденции ПРИКАЗЫВАЮ:", "justify", False, 12),
    ("1. Ввести в опытную эксплуатацию сервис DocMat для распознавания и конвертации "
     "электронных докуметов.", "justify", False, 12),
    ("2. Назначить ответственным за эксплуатацию сервиса начальника отдела информационых "
     "технологий.", "justify", False, 12),
    ("3. Обеспечить обработку документов без сохранения их на сервере.", "justify", False, 12),
    ("4. Контроль за исполнением настоящего приказа оставляю за собой.", "justify", False, 12),
    ("Генеральный директор\tИ. И. Иванов", "tab", False, 12),
]


def build_docx() -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.left_margin, section.right_margin = Pt(85), Pt(42)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for text, align, bold, size in CONTENT:
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_after = Pt(10)
        if align == "tab":
            fmt.tab_stops.add_tab_stop(section.page_width - section.left_margin - section.right_margin,
                                       WD_TAB_ALIGNMENT.RIGHT)
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Pt(36)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "demo.docx"
        src.write_bytes(build_docx())
        pdf = (await office.convert(src, "pdf", Path(tmp) / "out")).read_bytes()

    with pymupdf.open(stream=pdf, filetype="pdf") as text_pdf, pymupdf.open() as scan:
        for page in text_pdf:
            png = page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY).tobytes("png")
            new = scan.new_page(width=page.rect.width, height=page.rect.height)
            new.insert_image(new.rect, stream=png)
        OUT.parent.mkdir(parents=True, exist_ok=True)
        scan.save(OUT, deflate=True)
    print(OUT)


if __name__ == "__main__":
    asyncio.run(main())
