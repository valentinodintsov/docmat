"""Текст базового документа и перенос в него правок без потери оформления.

Базовый документ — DOCX (текстовые документы, PDF, сканы) или XLSX (таблицы).
Его текст собирается из «кусков»: редактируемых (текст w:t в DOCX, значение ячейки
в XLSX) и служебных (переводы строк между абзацами, табуляции между ячейками).
Отредактированный пользователем текст сравнивается с исходным, и каждое изменение
записывается в те куски, на которые оно приходится. Остальная разметка — шрифты,
стили, таблицы, колонтитулы, картинки — остаётся нетронутой.
"""

import datetime as dt
import difflib
import io
import re
from dataclasses import dataclass, field

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

_TOKEN = re.compile(r"\w+|\s|[^\w\s]")
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
W_P, W_TBL, W_TR, W_TC, W_R = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc"), qn("w:r")
W_T, W_TAB, W_BR, W_CR, W_SDT, W_SDT_CONTENT = (
    qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:sdt"), qn("w:sdtContent"),
)


@dataclass
class Piece:
    text: str
    target: object | None = None  # None — служебный кусок, правки в него не попадают
    start: int = 0
    original: str = ""


@dataclass
class Pieces:
    items: list[Piece] = field(default_factory=list)

    def add(self, text: str, target: object | None = None) -> None:
        if text:
            self.items.append(Piece(text, target))

    def finish(self) -> str:
        pos = 0
        for p in self.items:
            p.start, p.original = pos, p.text
            pos += len(p.text)
        return "".join(p.text for p in self.items)


@dataclass
class ApplyResult:
    data: bytes
    applied: int
    skipped: int  # правки, затрагивающие только служебные символы (например, склейка абзацев)


# ---------- DOCX ----------


def _in_fallback(el) -> bool:
    parent = el.getparent()
    while parent is not None:
        if parent.tag == _MC_FALLBACK:
            return True
        parent = parent.getparent()
    return False


def _paragraph(p_el, out: Pieces) -> None:
    for el in p_el.iter(W_T, W_TAB, W_BR, W_CR):
        if el.getparent().tag != W_R or _in_fallback(el):
            continue  # w:tab внутри w:tabs — это позиция табуляции, а не символ
        if el.tag == W_T:
            out.add(el.text or "", el)
        elif el.tag == W_TAB:
            out.add("\t")
        else:
            out.add("\n")


def _blocks(parent, out: Pieces, sep: str) -> None:
    first = True
    for child in parent.iterchildren(W_P, W_TBL, W_SDT):
        if child.tag == W_SDT:
            content = child.find(W_SDT_CONTENT)
            if content is not None:
                if not first:
                    out.add(sep)
                _blocks(content, out, sep)
                first = False
            continue
        if not first:
            out.add(sep)
        first = False
        if child.tag == W_P:
            _paragraph(child, out)
        else:
            _table(child, out)


def _table(tbl_el, out: Pieces) -> None:
    for r, tr in enumerate(tbl_el.iterchildren(W_TR)):
        if r:
            out.add("\n")
        for c, tc in enumerate(tr.iterchildren(W_TC)):
            if c:
                out.add("\t")
            _blocks(tc, out, " ")


def docx_pieces(doc) -> tuple[str, Pieces]:
    out = Pieces()
    _blocks(doc.element.body, out, "\n")
    return out.finish(), out


def _set_docx_text(t_el, text: str) -> None:
    t_el.text = text
    if text != text.strip():
        t_el.set(qn("xml:space"), "preserve")


# ---------- XLSX ----------


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.strftime("%d.%m.%Y %H:%M") if value.time() != dt.time() else value.strftime("%d.%m.%Y")
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    return str(value).replace("\t", " ").replace("\n", " ")


@dataclass
class _CellRef:
    cell: object
    numeric: bool


def xlsx_pieces(wb, wb_values) -> tuple[str, Pieces]:
    out = Pieces()
    first_sheet = True
    for ws, ws_values in zip(wb.worksheets, wb_values.worksheets, strict=True):
        rows = []
        for row, row_values in zip(ws.iter_rows(), ws_values.iter_rows(), strict=True):
            cells = list(zip(row, row_values, strict=True))
            while cells and cell_text(cells[-1][1].value) == "":
                cells.pop()
            rows.append(cells)
        while rows and not rows[-1]:
            rows.pop()
        if not rows:
            continue
        if not first_sheet:
            out.add("\n\n")
        first_sheet = False
        for r, cells in enumerate(rows):
            if r:
                out.add("\n")
            for c, (cell, value_cell) in enumerate(cells):
                if c:
                    out.add("\t")
                text = cell_text(value_cell.value)
                editable = (
                    not isinstance(cell, MergedCell)
                    and cell.data_type in ("s", "n", "inlineStr")
                    and not (isinstance(cell.value, str) and cell.value.startswith("="))
                )
                ref = _CellRef(cell, isinstance(cell.value, int | float)) if editable else None
                out.add(text, ref)
                if editable and not text:
                    out.items.append(Piece("", ref))  # пустая ячейка тоже может получить текст
    return out.finish(), out


def _number(text: str) -> int | float | None:
    s = text.strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return int(s) if re.fullmatch(r"-?\d+", s) else float(s)
    except ValueError:
        return None


def _set_xlsx_text(ref: _CellRef, text: str) -> None:
    if not text.strip():
        ref.cell.value = None
    elif ref.numeric and (num := _number(text)) is not None:
        ref.cell.value = num
    else:
        ref.cell.value = text


# ---------- Общая часть ----------


def _tokens(text: str) -> tuple[list[str], list[int]]:
    tokens, offsets = [], []
    for m in _TOKEN.finditer(text):
        tokens.append(m.group())
        offsets.append(m.start())
    offsets.append(len(text))
    return tokens, offsets


_LAYOUT_SPLIT = re.compile(r"([\t\n])")


def _split_by_layout(start: int, old: str, new: str) -> list[tuple[int, int, str]]:
    """Правка через границу абзаца или ячейки делится на правки внутри каждого куска.

    Иначе исправление двух соседних слов «Сокращеныя⏎времеми» превратилось бы в
    одну правку, и перевод строки оказался бы не на своём месте.
    """
    old_parts, new_parts = _LAYOUT_SPLIT.split(old), _LAYOUT_SPLIT.split(new)
    if len(old_parts) == 1 or old_parts[1::2] != new_parts[1::2]:
        return [(start, start + len(old), new)]
    edits, pos = [], start
    for o, n in zip(old_parts, new_parts, strict=True):
        if o != n:
            edits.append((pos, pos + len(o), n))
        pos += len(o)
    return edits


def diff(old: str, new: str) -> list[tuple[int, int, str]]:
    """Правки (начало, конец, новый текст) в координатах старого текста, по словам."""
    a, a_off = _tokens(old)
    b, b_off = _tokens(new)
    matcher = difflib.SequenceMatcher(None, a, b, autojunk=False)
    edits = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag != "equal":
            start, end = a_off[i1], a_off[i2]
            edits.extend(_split_by_layout(start, old[start:end], new[b_off[j1] : b_off[j2]]))
    return edits


def _owner_for_insert(pieces: list[Piece], pos: int) -> Piece | None:
    editable = [p for p in pieces if p.target is not None]
    for p in editable:  # предпочитаем кусок, который заканчивается в точке вставки
        if p.start < pos <= p.start + len(p.original) or (p.start == pos and not p.original):
            return p
    for p in editable:
        if p.start <= pos < p.start + len(p.original):
            return p
    left = [p for p in editable if p.start + len(p.original) <= pos]
    return left[-1] if left else (editable[0] if editable else None)


def apply_edits(pieces: Pieces, edits: list[tuple[int, int, str]], single_line: bool) -> tuple[int, int]:
    applied = skipped = 0
    for start, end, new in reversed(edits):
        if single_line:
            new = new.replace("\n", " ").replace("\t", " ")
        overlapped = [
            p for p in pieces.items
            if p.target is not None and p.start < end and p.start + len(p.original) > start
        ]
        if start == end or not overlapped:
            owner = _owner_for_insert(pieces.items, start)
            if owner is None or (not overlapped and start != end and not new):
                skipped += 1
                continue
            overlapped = [owner]
        owner = overlapped[0]
        # Удаляем справа налево, чтобы локальные смещения оставались верными.
        for p in reversed(overlapped):
            lo = max(start, p.start) - p.start
            hi = min(end, p.start + len(p.original)) - p.start
            if p is owner:
                p.text = p.text[:lo] + new + p.text[max(hi, lo) :]
            else:
                p.text = p.text[:lo] + p.text[hi:]
        applied += 1
    return applied, skipped


def text_of(kind: str, data: bytes) -> str:
    if kind == "docx":
        return docx_pieces(DocxDocument(io.BytesIO(data)))[0]
    wb = load_workbook(io.BytesIO(data))
    wb_values = load_workbook(io.BytesIO(data), data_only=True)
    return xlsx_pieces(wb, wb_values)[0]


def apply_text(kind: str, data: bytes, new_text: str) -> ApplyResult:
    """Переносит отредактированный текст в базовый документ."""
    new_text = new_text.replace("\r\n", "\n")
    if kind == "docx":
        doc = DocxDocument(io.BytesIO(data))
        old, pieces = docx_pieces(doc)
        if old == new_text:
            return ApplyResult(data, 0, 0)
        applied, skipped = apply_edits(pieces, diff(old, new_text), single_line=True)
        for p in pieces.items:
            if p.target is not None and p.text != p.original:
                _set_docx_text(p.target, p.text)
        buf = io.BytesIO()
        doc.save(buf)
        return ApplyResult(buf.getvalue(), applied, skipped)

    wb = load_workbook(io.BytesIO(data))
    wb_values = load_workbook(io.BytesIO(data), data_only=True)
    old, pieces = xlsx_pieces(wb, wb_values)
    if old == new_text:
        return ApplyResult(data, 0, 0)
    applied, skipped = apply_edits(pieces, diff(old, new_text), single_line=False)
    changed: dict[int, tuple[_CellRef, str]] = {}
    for p in pieces.items:
        if p.target is not None:
            prev = changed.get(id(p.target), (p.target, ""))[1]
            changed[id(p.target)] = (p.target, prev + p.text)
    originals: dict[int, str] = {}
    for p in pieces.items:
        if p.target is not None:
            originals[id(p.target)] = originals.get(id(p.target), "") + p.original
    for key, (ref, text) in changed.items():
        if text != originals[key]:
            _set_xlsx_text(ref, text)
    buf = io.BytesIO()
    wb.save(buf)
    return ApplyResult(buf.getvalue(), applied, skipped)
