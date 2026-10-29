"""Восстановление вёрстки PDF и сканов в редактируемый DOCX.

Источник — текстовый слой PDF (шрифт, кегль, начертание, цвет каждого фрагмента)
или слова, найденные Tesseract (только координаты, кегль оценивается по высоте строки).
Из координат восстанавливаются: размер страницы и поля, абзацы, выравнивание,
отступы, межстрочный интервал, табуляции (текст «слева … справа» в одной строке),
таблицы и изображения (для PDF с текстовым слоем).
"""

import io
import re
import statistics
from dataclasses import dataclass, field

import numpy as np
import pymupdf
import pytesseract
from docx import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor
from PIL import Image, ImageFilter, ImageOps

from app.config import settings
from app.services import spellcheck

# Для сканов шрифт неизвестен; в российском делопроизводстве обычно Times New Roman.
OCR_FONT = "Times New Roman"
# Отношение кегля к высоте строки, найденной Tesseract (от верха прописных до низа
# выносных элементов) — запасная оценка, если в строке нет слов-эталонов.
OCR_SIZE_FACTOR = 1.15
COMMON_SIZES = (8, 9, 10, 10.5, 11, 12, 13, 14, 15, 16, 18, 20, 22, 24, 26, 28, 32, 36, 40, 48)
FONT_NAMES = {
    "timesnewroman": "Times New Roman", "times": "Times New Roman", "arial": "Arial", "helvetica": "Arial",
    "calibri": "Calibri", "cambria": "Cambria", "couriernew": "Courier New", "courier": "Courier New",
    "liberationserif": "Liberation Serif", "liberationsans": "Liberation Sans", "liberationmono": "Liberation Mono",
    "dejavusans": "DejaVu Sans", "dejavuserif": "DejaVu Serif", "dejavusansmono": "DejaVu Sans Mono",
    "ptsans": "PT Sans", "ptserif": "PT Serif", "ptastrasans": "PT Astra Sans", "ptastraserif": "PT Astra Serif",
    "verdana": "Verdana", "tahoma": "Tahoma", "georgia": "Georgia", "segoeui": "Segoe UI",
    "carlito": "Carlito", "caladea": "Caladea", "opensans": "Open Sans", "roboto": "Roboto",
}
_FONT_SUFFIX = re.compile(r"[-,](Bold|Italic|Oblique|Regular|Roman|Book|Medium|Light|Semibold|BoldItalic)\w*$", re.I)


@dataclass
class Span:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    bold: bool = False
    italic: bool = False
    color: str | None = None
    font: str | None = None


@dataclass
class Chunk:
    """Строка из источника (строка PyMuPDF или строка Tesseract)."""

    spans: list[Span]
    group: int  # блок источника: строки разных блоков не объединяются в один абзац

    @property
    def y0(self) -> float:
        return min(s.y0 for s in self.spans)

    @property
    def y1(self) -> float:
        return max(s.y1 for s in self.spans)

    @property
    def x0(self) -> float:
        return min(s.x0 for s in self.spans)

    @property
    def x1(self) -> float:
        return max(s.x1 for s in self.spans)


@dataclass
class Line:
    spans: list[Span]
    group: int
    x0: float = 0
    y0: float = 0
    x1: float = 0
    y1: float = 0
    size: float = 12
    tabs: list[tuple[float, str]] = field(default_factory=list)  # (позиция, "left"/"right")

    def __post_init__(self):
        self.spans.sort(key=lambda s: s.x0)
        self.x0 = min(s.x0 for s in self.spans)
        self.x1 = max(s.x1 for s in self.spans)
        self.y0 = min(s.y0 for s in self.spans)
        self.y1 = max(s.y1 for s in self.spans)
        self.size = max(self.spans, key=lambda s: len(s.text.strip())).size

    @property
    def height(self) -> float:
        return self.y1 - self.y0


@dataclass
class Para:
    lines: list[Line]
    align: str = "left"
    left_indent: float = 0
    first_indent: float = 0
    line_pitch: float | None = None

    @property
    def y0(self) -> float:
        return self.lines[0].y0

    @property
    def y1(self) -> float:
        return self.lines[-1].y1


@dataclass
class TableBlock:
    y0: float
    y1: float
    x0: float
    x1: float
    rows: list[list[str]]
    col_widths: list[float]
    size: float = 11
    font: str | None = None


@dataclass
class ImageBlock:
    y0: float
    y1: float
    x0: float
    x1: float
    png: bytes


@dataclass
class PageLayout:
    width: float
    height: float
    items: list = field(default_factory=list)
    left: float = 56
    right: float = 56


# ---------- Источник: текстовый слой PDF ----------


def font_family(name: str | None) -> str | None:
    if not name:
        return None
    name = name.split("+", 1)[-1]  # префикс подмножества «ABCDEF+»
    name = _FONT_SUFFIX.sub("", name)
    name = re.sub(r"(PSMT|PS|MT)$", "", name)
    key = re.sub(r"[^a-z]", "", name.lower())
    if key in FONT_NAMES:
        return FONT_NAMES[key]
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name).strip() or None


def _color(value: int) -> str | None:
    return None if value in (0, None) else f"{value:06X}"


def _inside(item, box) -> bool:
    return _inside_box((item.x0, item.y0, item.x1, item.y1), box)


def _tables(page: pymupdf.Page) -> list[TableBlock]:
    try:
        found = page.find_tables()
    except Exception:  # поиск таблиц — эвристика PyMuPDF, на сложных страницах он может падать
        return []
    words = page.get_text("words")
    spans = [s for b in page.get_text("dict")["blocks"] if b.get("type") == 0 for ln in b["lines"] for s in ln["spans"]]
    tables = []
    for t in found.tables:
        if not t.rows or len(t.rows[0].cells) < 2:
            continue
        # Текст ячейки собирается из слов: extract() на некоторых PDF теряет пробелы.
        rows = []
        for row in t.rows:
            cells = []
            for cell in row.cells:
                if cell is None:
                    cells.append("")
                    continue
                cells.append(_join_words([w for w in words if _inside_box(w[:4], cell)]))
            rows.append(cells)
        x0, y0, x1, y1 = t.bbox
        n = len(t.rows[0].cells)
        widths = [(c[2] - c[0]) if c else (x1 - x0) / n for c in t.rows[0].cells]
        inner = [s for s in spans if _inside_box(s["bbox"], t.bbox) and s["text"].strip()]
        size = statistics.median(round(s["size"] * 2) / 2 for s in inner) if inner else 11
        font = font_family(inner[0]["font"]) if inner else None
        tables.append(TableBlock(y0, y1, x0, x1, rows, widths, size, font))
    return tables


def _join_words(words: list) -> str:
    """Слова в порядке чтения. Символы другим шрифтом («А» + «4») PDF кладёт
    отдельной строкой, поэтому строки определяются по геометрии, а слова без
    зазора склеиваются без пробела."""
    rows: list[list] = []
    for w in sorted(words, key=lambda w: (w[1] + w[3]) / 2):
        cy, h = (w[1] + w[3]) / 2, w[3] - w[1]
        if rows and abs(cy - (rows[-1][0][1] + rows[-1][0][3]) / 2) < 0.5 * h:
            rows[-1].append(w)
        else:
            rows.append([w])
    out = []
    for row in rows:
        row.sort(key=lambda w: w[0])
        text = row[0][4]
        for prev, w in zip(row, row[1:]):
            text += ("" if w[0] - prev[2] < 0.15 * (w[3] - w[1]) else " ") + w[4]
        out.append(text)
    return " ".join(out)


def _inside_box(bbox, box) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return box[0] <= cx <= box[2] and box[1] <= cy <= box[3]


def _images(page: pymupdf.Page, blocks: list[dict]) -> list[ImageBlock]:
    images = []
    page_area = page.rect.width * page.rect.height
    for b in blocks:
        if b.get("type") != 1:
            continue
        x0, y0, x1, y1 = b["bbox"]
        if (x1 - x0) * (y1 - y0) > 0.6 * page_area or (x1 - x0) < 8 or (y1 - y0) < 8:
            continue  # фон-скан под текстовым слоем или мелкий мусор
        try:
            png = pymupdf.Pixmap(b["image"]).tobytes("png")
        except Exception:  # экзотические кодеки (JBIG2, JPX) — изображение пропускается
            continue
        images.append(ImageBlock(y0, y1, x0, x1, png))
    return images


def from_text_layer(page: pymupdf.Page) -> PageLayout:
    data = page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT | pymupdf.TEXT_PRESERVE_IMAGES)
    tables = _tables(page)
    table_boxes = [(t.x0, t.y0, t.x1, t.y1) for t in tables]
    chunks = []
    for b_no, block in enumerate(data["blocks"]):
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            if abs(line["dir"][1]) > 0.1:
                continue  # повёрнутый текст
            spans = [
                Span(
                    s["text"], *s["bbox"], size=round(s["size"] * 2) / 2,
                    bold=bool(s["flags"] & pymupdf.TEXT_FONT_BOLD) or "bold" in s["font"].lower(),
                    italic=bool(s["flags"] & pymupdf.TEXT_FONT_ITALIC) or "italic" in s["font"].lower(),
                    color=_color(s["color"]), font=font_family(s["font"]),
                )
                for s in line["spans"]
                if s["text"].strip()
            ]
            chunk = Chunk(spans, b_no) if spans else None
            if chunk and not any(_inside(chunk, box) for box in table_boxes):
                chunks.append(chunk)
    # Блоки PyMuPDF не совпадают с абзацами (абзац по ширине часто разбит на несколько),
    # поэтому для текстового слоя абзацы определяются только по геометрии.
    layout = build_page(page.rect.width, page.rect.height, chunks, use_groups=False)
    layout.items.extend(tables)
    layout.items.extend(_images(page, data["blocks"]))
    layout.items.sort(key=lambda it: it.y0)
    _update_margins(layout)
    return layout


# ---------- Источник: OCR ----------

# Слова только из строчных букв без выносных элементов имеют высоту строчных (x-height),
# слова из прописных и цифр — высоту прописных.
_X_HEIGHT_WORD = re.compile(r"[авгежзиклмнопстхчшъыьэюяacemnorsuvwxz]+")
_CAP_HEIGHT_WORD = re.compile(r"[АБВГЕЖЗИКЛМНОПРСТУФХЧШЪЫЬЭЮЯA-IK-PR-Z0-9]+")
# Доли кегля по измерениям Tesseract на сканах 200–300 dpi (Times New Roman 12–16 pt).
X_HEIGHT_EM = 0.48
CAP_HEIGHT_EM = 0.68
# Строка жирная, если доля тёмных пикселей в словах заметно выше типичной для страницы.
BOLD_INK_RATIO = 1.2


def _snap_size(size: float) -> float:
    return min(COMMON_SIZES, key=lambda s: abs(s - size))


def _em_px(word: str, height: int) -> float | None:
    if _X_HEIGHT_WORD.fullmatch(word):
        return height / X_HEIGHT_EM
    if _CAP_HEIGHT_WORD.fullmatch(word):
        return height / CAP_HEIGHT_EM
    return None


def _otsu(gray: np.ndarray) -> int:
    hist = np.bincount(gray.ravel(), minlength=256).astype(float)
    p = hist / hist.sum()
    w = np.cumsum(p)
    mu = np.cumsum(p * np.arange(256))
    return int(np.argmax((mu[-1] * w - mu) ** 2 / (w * (1 - w) + 1e-12)))


def _binarize(image: Image.Image) -> Image.Image:
    gray = np.asarray(ImageOps.autocontrast(image))
    return Image.fromarray(((gray > _otsu(gray)) * 255).astype(np.uint8))


# Варианты очистки для зашумлённых и низкокачественных сканов (подобраны в
# scripts/eval_correction.py). Применяются, только если исходное изображение
# распознано плохо; побеждает вариант с наибольшим числом словарных слов.
_CLEANUPS = (
    lambda img: img.filter(ImageFilter.MedianFilter(5)),
    lambda img: img.filter(ImageFilter.MedianFilter(7)),
    lambda img: _binarize(img.filter(ImageFilter.GaussianBlur(1.6))),
)
GOOD_OCR_RATIO = 0.85


def _tesseract(image: Image.Image, lang: str) -> dict:
    return pytesseract.image_to_data(image, lang=lang, config="--psm 3", output_type=pytesseract.Output.DICT)


def _dictionary_words(data: dict) -> tuple[int, int]:
    """(слов из словаря, всего слов) — мера качества распознавания без эталона."""
    text = " ".join(t for t in data["text"] if t.strip())
    words = [m.group() for m in spellcheck.WORD_RE.finditer(text) if len(m.group()) >= 3]
    return sum(spellcheck.check_word(w) for w in words), len(words)


def recognize_image(image: Image.Image, lang: str) -> dict:
    data = _tesseract(image, lang)
    valid, total = _dictionary_words(data)
    if total == 0 or valid / total >= GOOD_OCR_RATIO:
        return data
    best, best_valid = data, valid
    for cleanup in _CLEANUPS:
        candidate = _tesseract(cleanup(image), lang)
        cand_valid, _ = _dictionary_words(candidate)
        if cand_valid > best_valid:
            best, best_valid = candidate, cand_valid
    return best


def from_ocr(page: pymupdf.Page, lang: str) -> PageLayout:
    dpi = settings.ocr_dpi
    pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    ink = np.asarray(image) < 128
    data = recognize_image(image, lang)
    k = 72 / dpi
    lines: dict[tuple[int, int, int], list[int]] = {}
    for i, word in enumerate(data["text"]):
        if data["level"][i] == 5 and word.strip():
            lines.setdefault((data["block_num"][i], data["par_num"][i], data["line_num"][i]), []).append(i)
    if not lines:
        return build_page(page.rect.width, page.rect.height, [])

    stats = {}
    for key, idx in lines.items():
        top = min(data["top"][i] for i in idx)
        bottom = max(data["top"][i] + data["height"][i] for i in idx)
        ems = [e for i in idx if (e := _em_px(data["text"][i], data["height"][i]))]
        size = statistics.median(ems) * k if ems else None
        dark = area = 0
        for i in idx:
            x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
            dark += int(ink[y : y + h, x : x + w].sum())
            area += w * h
        stats[key] = (top * k, bottom * k, size, dark / max(area, 1))

    # Кегль основного текста: оценки строк рядом с ним приравниваются к нему,
    # чтобы шум измерений не давал «пёстрых» размеров в одном абзаце.
    measured = [v[2] for v in stats.values() if v[2]]
    typical_height = statistics.median(v[1] - v[0] for v in stats.values())
    body = _snap_size(statistics.median(measured)) if measured else _snap_size(typical_height * OCR_SIZE_FACTOR)
    typical_ink = statistics.median(v[3] for v in stats.values())
    chunks = []
    for (block, par, line_no), idx in lines.items():
        top, bottom, size, density = stats[(block, par, line_no)]
        if size is None:
            # В строке нет слов-эталонов. Высота строки зависит от набора букв, поэтому
            # по ней судим только о заметно более крупном тексте, иначе — основной кегль.
            ratio = (bottom - top) / typical_height
            size = body * ratio if ratio > 1.3 else body
        size = body if abs(size - body) < 0.15 * body else _snap_size(size)
        bold = density > BOLD_INK_RATIO * typical_ink
        spans = [
            Span(
                data["text"][i], data["left"][i] * k, top,
                (data["left"][i] + data["width"][i]) * k, bottom, size=size, bold=bold, font=OCR_FONT,
            )
            for i in idx
        ]
        chunks.append(Chunk(spans, block * 1000 + par))
    layout = build_page(page.rect.width, page.rect.height, chunks)
    _update_margins(layout)
    return layout


# ---------- Восстановление абзацев ----------


def _vertical_overlap(a, b) -> float:
    inter = min(a.y1, b.y1) - max(a.y0, b.y0)
    return inter / max(min(a.y1 - a.y0, b.y1 - b.y0), 0.1)


def _merge_visual_lines(chunks: list[Chunk]) -> list[Line]:
    """Объединяет фрагменты, стоящие на одной строке (например, «г. Москва … 15 октября»)."""
    rows: list[list[Chunk]] = []
    def disjoint(row: list[Chunk], chunk: Chunk) -> bool:
        # Сравниваются отдельные фрагменты: знаки препинания PDF нередко лежат
        # отдельной «строкой» внутри габаритов основной строки.
        tol = 0.5
        return all(a.x1 <= b.x0 + tol or a.x0 >= b.x1 - tol for c in row for a in c.spans for b in chunk.spans)

    for chunk in sorted(chunks, key=lambda c: (c.y0, c.x0)):
        for row in rows:
            if any(_vertical_overlap(c, chunk) > 0.5 for c in row) and disjoint(row, chunk):
                row.append(chunk)
                break
        else:
            rows.append([chunk])
    lines = []
    for row in rows:
        row.sort(key=lambda c: c.x0)
        lines.append(Line([s for c in row for s in c.spans], row[0].group))
    lines.sort(key=lambda ln: ln.y0)
    return lines


def _find_tabs(line: Line, left: float, right: float) -> None:
    """Большой разрыв между словами — это табуляция, а не пробелы."""
    tol = max(line.size, 6)
    segments = [[line.spans[0]]]
    for prev, cur in zip(line.spans, line.spans[1:]):
        if cur.x0 - prev.x1 > 2.5 * line.size:
            segments.append([cur])
        else:
            segments[-1].append(cur)
    if len(segments) < 2:
        return
    for seg in segments[1:]:
        seg_x0, seg_x1 = seg[0].x0, max(s.x1 for s in seg)
        if right - seg_x1 < tol and seg is segments[-1]:
            line.tabs.append((right - left, "right"))
        else:
            line.tabs.append((seg_x0 - left, "left"))
        seg[0].text = "\t" + seg[0].text.lstrip()


def _word_gap(line: Line) -> float | None:
    """Средний пробел между словами в долях кегля (только для OCR, где фрагмент — слово)."""
    gaps = [b.x0 - a.x1 for a, b in zip(line.spans, line.spans[1:]) if not b.text.startswith("\t")]
    return statistics.mean(gaps) / line.size if len(gaps) >= 2 else None


def _alignment(para: Para, left: float, right: float, page_gap: float | None) -> None:
    tol = max(para.lines[0].size, 8)
    width = right - left
    gaps = [(ln.x0 - left, right - ln.x1) for ln in para.lines]
    body = gaps[:-1] if len(gaps) > 1 else gaps
    if para.lines[0].tabs:
        para.align = "left"
    elif all(abs(lg - rg) < tol and lg > tol for lg, rg in gaps):
        para.align = "center"
    elif all(rg < tol and lg > 0.25 * width for lg, rg in gaps):
        para.align = "right"
    # По ширине — только если до правого края доходят хотя бы две строки (одна
    # самая длинная строка доходит до края по определению).
    elif len(body) > 1 and all(rg < tol * 0.6 for _, rg in body):
        para.align = "justify"
    elif (
        page_gap and len(gaps) > 1 and all(rg < tol * 0.6 for _, rg in body)
        and all((g := _word_gap(ln)) and g > 1.25 * page_gap for ln in para.lines[:-1])
    ):
        para.align = "justify"  # пробелы растянуты, чтобы строка дошла до края
    else:
        para.align = "left"
    if para.align in ("left", "justify"):
        rest = [lg for lg, _ in gaps[1:]] or [gaps[0][0]]
        para.left_indent = max(min(rest), 0)
        para.first_indent = gaps[0][0] - para.left_indent
    if len(para.lines) > 1:
        para.line_pitch = statistics.median(b.y0 - a.y0 for a, b in zip(para.lines, para.lines[1:]))


def _paragraphs(lines: list[Line], left: float, right: float, use_groups: bool) -> list[Para]:
    if not lines:
        return []
    median_h = statistics.median(ln.height for ln in lines)
    paras: list[Para] = []
    for line in lines:
        cur = paras[-1] if paras else None
        prev = cur.lines[-1] if cur else None
        new = (
            cur is None
            or (use_groups and line.group != prev.group)
            or line.y0 - prev.y1 > 0.7 * median_h
            or abs(line.size - prev.size) > 1.5
            or bool(line.tabs) or bool(prev.tabs)
            # Строка, не дошедшая до правого края, заканчивает абзац.
            or right - prev.x1 > max(0.15 * (right - left), 4 * prev.size)
        )
        if new:
            paras.append(Para([line]))
        else:
            cur.lines.append(line)
    word_gaps = [g for ln in lines if (g := _word_gap(ln))]
    page_gap = statistics.median(word_gaps) if word_gaps else None
    for p in paras:
        _alignment(p, left, right, page_gap)
    # Если на странице есть явно выровненные по ширине абзацы, то и двухстрочные
    # абзацы, первая строка которых доходит до края, выровнены так же.
    if any(p.align == "justify" for p in paras):
        for p in paras:
            if p.align == "left" and len(p.lines) == 2 and right - p.lines[0].x1 < 0.6 * max(p.lines[0].size, 8):
                p.align = "justify"
    return paras


def build_page(width: float, height: float, chunks: list[Chunk], use_groups: bool = True) -> PageLayout:
    layout = PageLayout(width, height, left=56, right=width - 56)
    if not chunks:
        return layout
    lines = _merge_visual_lines(chunks)
    left = min(ln.x0 for ln in lines)
    right = max(ln.x1 for ln in lines)
    for ln in lines:
        _find_tabs(ln, left, right)
    layout.left, layout.right = left, right
    layout.items = _paragraphs(lines, left, right, use_groups)
    return layout


def _update_margins(layout: PageLayout) -> None:
    """Таблицы и изображения могут выходить за границы текста — расширяем область."""
    boxes = [it for it in layout.items if not isinstance(it, Para)]
    if boxes:
        layout.left = min([layout.left] + [b.x0 for b in boxes])
        layout.right = max([layout.right] + [b.x1 for b in boxes])


# ---------- Запись DOCX ----------

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _line_text_spans(para: Para) -> list[Span]:
    """Строки абзаца в один поток фрагментов: пробелы между словами, склейка переносов."""
    out: list[Span] = []
    for n, line in enumerate(para.lines):
        for i, span in enumerate(line.spans):
            text = span.text
            if out:
                prev = out[-1]
                if i == 0:  # стык строк
                    if re.search(r"\w-$", prev.text) and text[:1].islower():
                        prev.text = prev.text[:-1]
                    elif not prev.text.endswith(" ") and not text.startswith(" "):
                        text = " " + text
                elif not text.startswith(("\t", " ")) and not prev.text.endswith(" ") and span.x0 - prev.x1 > 0.12 * span.size:
                    text = " " + text
            out.append(Span(text, span.x0, span.y0, span.x1, span.y1, span.size, span.bold, span.italic, span.color, span.font))
    merged: list[Span] = []
    for s in out:
        if merged and (merged[-1].size, merged[-1].bold, merged[-1].italic, merged[-1].color, merged[-1].font) == (
            s.size, s.bold, s.italic, s.color, s.font
        ):
            merged[-1].text += s.text
        else:
            merged.append(s)
    if merged:
        merged[-1].text = merged[-1].text.rstrip()
    return merged


def _write_para(doc, para: Para, space_before: float) -> None:
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.alignment = _ALIGN[para.align]
    fmt.space_before = Pt(max(round(space_before), 0))
    fmt.space_after = Pt(0)
    if para.left_indent > 1:
        fmt.left_indent = Pt(round(para.left_indent))
    if abs(para.first_indent) > 1:
        fmt.first_line_indent = Pt(round(para.first_indent))
    if para.line_pitch:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(round(para.line_pitch, 1))
    for pos, kind in para.lines[0].tabs:
        fmt.tab_stops.add_tab_stop(
            Pt(round(pos)), WD_TAB_ALIGNMENT.RIGHT if kind == "right" else WD_TAB_ALIGNMENT.LEFT
        )
    for span in _line_text_spans(para):
        run = p.add_run(span.text)
        run.font.size = Pt(span.size)
        run.bold = span.bold or None
        run.italic = span.italic or None
        if span.font:
            run.font.name = span.font
        if span.color:
            run.font.color.rgb = RGBColor.from_string(span.color)


def _write_table(doc, table: TableBlock) -> None:
    cols = max(len(r) for r in table.rows)
    t = doc.add_table(rows=len(table.rows), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(table.rows):
        for c in range(cols):
            cell = t.cell(r, c)
            run = cell.paragraphs[0].add_run(row[c] if c < len(row) else "")
            run.font.size = Pt(table.size)
            if table.font:
                run.font.name = table.font
            if c < len(table.col_widths):
                cell.width = Pt(table.col_widths[c])


def _write_image(doc, image: ImageBlock, left: float, right: float) -> None:
    p = doc.add_paragraph()
    width = image.x1 - image.x0
    center = (image.x0 + image.x1) / 2
    if abs(center - (left + right) / 2) < 0.1 * (right - left):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif image.x0 - left > 0.4 * (right - left):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(io.BytesIO(image.png), width=Pt(min(width, right - left)))


def to_docx(pages: list[PageLayout]) -> bytes:
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = OCR_FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for n, page in enumerate(pages):
        section = doc.sections[0] if n == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width, section.page_height = Pt(page.width), Pt(page.height)
        left = max(page.left, 20)
        # Небольшой запас ширины: метрики подставленного шрифта могут чуть отличаться,
        # и самая длинная строка не должна переноситься.
        right_margin = max(page.width - page.right - 6, 20)
        top = min(max(page.items[0].y0, 20), 56) if page.items else 56
        section.left_margin, section.right_margin = Pt(round(left)), Pt(round(right_margin))
        section.top_margin, section.bottom_margin = Pt(round(top)), Pt(28)
        section.header_distance = section.footer_distance = Pt(14)

        prev_bottom = top
        for item in page.items:
            gap = item.y0 - prev_bottom
            if isinstance(item, Para):
                # Строка в Word выше видимых букв: часть зазора уже даёт сам межстрочный интервал.
                line = item.lines[0]
                _write_para(doc, item, gap - max(1.15 * line.size - line.height, 0))
            elif isinstance(item, TableBlock):
                _write_table(doc, item)
            else:
                _write_image(doc, item, page.left, page.right)
            prev_bottom = item.y1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
