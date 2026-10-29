import asyncio
import io
import tempfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.main import app
from app.services import office


@pytest.fixture(scope="session")
def client() -> TestClient:
    return TestClient(app)


def make_docx() -> bytes:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Техническое задание")
    run.bold = True
    run.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Сервис выполняет ")
    bold = p.add_run("оптическое распознавание")
    bold.bold = True
    p.add_run(" текста и проверку орфографии документов, переданных пользователем через веб-интерфейс.")
    doc.add_paragraph("Документы не хранятся на сервере после обработки.")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for r, row in enumerate([["Параметр", "Значение"], ["Формат", "PDF/A"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["Наименование", "Количество", "Цена", "Сумма"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.append(["Бумага", 10, 250, "=B2*C2"])
    ws.append(["Картридж", 2, 3500.5, "=B3*C3"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_docx() -> bytes:
    return make_docx()


@pytest.fixture(scope="session")
def sample_xlsx() -> bytes:
    # Через LibreOffice, чтобы в файле были вычисленные значения формул, как у файлов из Excel.
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "smeta.xlsx"
        src.write_bytes(make_xlsx())
        return asyncio.run(office.convert(src, "xlsx", Path(tmp) / "out")).read_bytes()


@pytest.fixture(scope="session")
def sample_pdf(sample_docx) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sample.docx"
        src.write_bytes(sample_docx)
        return asyncio.run(office.convert(src, "pdf", Path(tmp) / "out")).read_bytes()


@pytest.fixture(scope="session")
def scanned_pdf(sample_pdf) -> bytes:
    """PDF без текстового слоя: страница вставлена как изображение."""
    with pymupdf.open(stream=sample_pdf, filetype="pdf") as src, pymupdf.open() as out:
        page = src[0]
        png = page.get_pixmap(dpi=200).tobytes("png")
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=png)
        return out.tobytes()
