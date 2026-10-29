import io
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services import layout

DEMO = Path(__file__).resolve().parents[1] / "docs" / "samples" / "demo-scan.pdf"


def paras(page):
    return [it for it in page.items if isinstance(it, layout.Para)]


def text(para):
    return "".join(s.text for ln in para.lines for s in ln.spans)


def test_text_layer_restores_formatting(sample_pdf):
    with pymupdf.open(stream=sample_pdf, filetype="pdf") as pdf:
        page = layout.from_text_layer(pdf[0])
    title, body, *_ = paras(page)
    assert title.align == "center" and title.lines[0].size == 16 and title.lines[0].spans[0].bold
    assert body.align == "justify" and len(body.lines) >= 2
    assert any(s.bold and "оптическое" in s.text for ln in body.lines for s in ln.spans)
    (table,) = [it for it in page.items if isinstance(it, layout.TableBlock)]
    assert table.rows == [["Параметр", "Значение"], ["Формат", "PDF/A"]]


def test_docx_written_from_layout(sample_pdf):
    with pymupdf.open(stream=sample_pdf, filetype="pdf") as pdf:
        data = layout.to_docx([layout.from_text_layer(pdf[0])])
    doc = Document(io.BytesIO(data))
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.paragraphs[0].runs[0].bold
    assert doc.tables[0].cell(1, 1).text == "PDF/A"


def test_ocr_layout_of_demo_scan():
    with pymupdf.open(DEMO) as pdf:
        page = layout.from_ocr(pdf[0], "rus")
    items = paras(page)
    org, order, city, heading, *body = items
    assert order.align == "center" and order.lines[0].size == 16 and order.lines[0].spans[0].bold
    assert org.lines[0].size == 14
    assert city.lines[0].tabs and city.lines[0].tabs[-1][1] == "right"  # «г. Москва … дата» справа
    assert heading.align == "center" and heading.lines[0].spans[0].bold
    assert body[0].align == "justify" and body[0].first_indent > 20 and body[0].lines[0].size == 12


def test_font_family_names():
    assert layout.font_family("ABCDEF+TimesNewRomanPS-BoldMT") == "Times New Roman"
    assert layout.font_family("LiberationSerif-Italic") == "Liberation Serif"
    assert layout.font_family("SomeCustomFont") == "Some Custom Font"
