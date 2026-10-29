import base64
import dataclasses
import io
import os

import pymupdf
import pytest
from docx import Document
from openpyxl import load_workbook

from app.config import settings
from app.services import corrector


def upload(client, path, name, data, **form):
    return client.post(path, files={"file": (name, data)}, data=form)


def test_index_page(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Одинцов Валентин Николаевич, 2026" in r.text
    for label in ("Распознать", "Исправить", "Очистить", "Конвертировать"):
        assert label in r.text


def test_rejects_unsupported_extension(client):
    r = upload(client, "/api/preview", "image.png", b"\x89PNG....")
    assert r.status_code == 415


def test_rejects_spoofed_extension(client):
    r = upload(client, "/api/preview", "fake.pdf", b"not a pdf at all")
    assert r.status_code == 415


def test_preview_pdf(client, sample_pdf):
    r = upload(client, "/api/preview", "sample.pdf", sample_pdf)
    assert r.status_code == 200
    body = r.json()
    assert body["format"] == "pdf" and body["page_count"] == 1
    assert body["pages"][0].startswith("data:image/png;base64,")
    assert r.headers["cache-control"] == "no-store"


def test_preview_docx(client, sample_docx):
    r = upload(client, "/api/preview", "sample.docx", sample_docx)
    assert r.status_code == 200
    assert r.json()["format"] == "docx" and r.json()["pages"]


def recognize(client, name, data, **form) -> dict:
    r = upload(client, "/api/recognize", name, data, **form)
    assert r.status_code == 200, r.text
    return r.json()


def convert(client, body: dict, text: str | None = None, filename="document.docx") -> dict:
    base = base64.b64decode(body["base"]["data"])
    r = client.post(
        "/api/convert",
        files={"base": ("base", base), "text": ("text.txt", (text if text is not None else body["text"]).encode())},
        data={"kind": body["base"]["kind"], "filename": filename},
    )
    assert r.status_code == 200, r.text
    out = r.json()
    out["by_format"] = {f["format"]: base64.b64decode(f["data"]) for f in out["files"]}
    return out


def test_recognize_docx_uses_original_as_base(client, sample_docx):
    body = recognize(client, "sample.docx", sample_docx)
    assert body["pages"] == [{"page": None, "method": "document"}]
    assert body["base"]["kind"] == "docx"
    assert base64.b64decode(body["base"]["data"]) == sample_docx
    assert "Сервис выполняет оптическое распознавание текста и проверку орфографии" in body["text"]


def test_recognize_docx_force_ocr_goes_through_pdf(client, sample_docx):
    body = recognize(client, "sample.docx", sample_docx, force_ocr="true")
    assert body["pages"][0]["method"] == "ocr"


def test_recognize_xlsx_keeps_cells(client, sample_xlsx):
    body = recognize(client, "smeta.xlsx", sample_xlsx)
    assert body["base"]["kind"] == "xlsx"
    assert body["text"].splitlines()[1] == "Бумага\t10\t250\t2500"
    assert body["errors"] == []


def test_recognize_text_layer(client, sample_pdf):
    body = recognize(client, "sample.pdf", sample_pdf)
    assert body["pages"] == [{"page": 1, "method": "text"}]
    assert "Документы не хранятся на сервере" in body["text"]
    assert body["errors"] == []


def test_recognize_scanned_with_ocr(client, scanned_pdf):
    body = recognize(client, "scan.pdf", scanned_pdf, lang="rus")
    assert body["pages"][0]["method"] == "ocr"
    assert "сервере" in body["text"]
    assert body["base"]["kind"] == "docx"


def test_convert_docx_keeps_original_formatting(client, sample_docx):
    body = recognize(client, "sample.docx", sample_docx)
    out = convert(client, body, body["text"].replace("Документы", "Файлы"), filename="отчёт.docx")
    assert set(out["by_format"]) == {"docx", "xlsx", "odt", "ods", "pdf", "pdfa"}
    assert out["applied"] == 1 and out["skipped"] == 0
    assert out["files"][0]["filename"] == "отчёт.docx"
    doc = Document(io.BytesIO(out["by_format"]["docx"]))
    assert doc.paragraphs[0].runs[0].bold and doc.paragraphs[0].runs[0].font.size.pt == 16
    assert doc.paragraphs[2].text == "Файлы не хранятся на сервере после обработки."
    assert doc.tables[0].cell(1, 1).text == "PDF/A"
    with pymupdf.open(stream=out["by_format"]["pdfa"], filetype="pdf") as pdf:
        assert "pdfaid:part" in pdf.get_xml_metadata()
        assert "Файлы не хранятся" in pdf[0].get_text()
    ws = load_workbook(io.BytesIO(out["by_format"]["xlsx"])).active
    assert ws["A1"].value == "Техническое задание" and ws["A1"].font.bold


def test_convert_spreadsheet_keeps_table(client, sample_xlsx):
    body = recognize(client, "smeta.xlsx", sample_xlsx)
    out = convert(client, body, body["text"].replace("\t10\t", "\t20\t"))
    ws = load_workbook(io.BytesIO(out["by_format"]["xlsx"]), data_only=True).active
    assert ws["B2"].value == 20 and ws["D2"].value == 5000  # формула пересчитана
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("DDEBF7")
    doc = Document(io.BytesIO(out["by_format"]["docx"]))
    assert doc.tables and doc.tables[0].cell(1, 0).text == "Бумага"
    with pymupdf.open(stream=out["by_format"]["pdfa"], filetype="pdf") as pdf:
        assert "pdfaid:part" in pdf.get_xml_metadata()


def test_convert_rejects_broken_base(client):
    r = client.post(
        "/api/convert",
        files={"base": ("base", b"PK\x03\x04garbage"), "text": ("t.txt", "текст".encode())},
        data={"kind": "docx"},
    )
    assert r.status_code == 415


def test_check_endpoint(client):
    r = client.post("/api/check", json={"text": "Пример доккумента"})
    assert [e["word"] for e in r.json()["errors"]] == ["доккумента"]


def test_correct_falls_back_to_dictionary(client, monkeypatch):
    monkeypatch.setattr(corrector, "settings", dataclasses.replace(settings, ollama_url="http://127.0.0.1:9"))
    r = client.post("/api/correct", json={"text": "Пример доккумента"})
    body = r.json()
    assert body["engine"] == "hunspell"
    assert body["text"] == "Пример документа"


def test_nothing_left_in_tmp_dir(client, sample_docx):
    tmp = settings.tmp_dir or "/tmp"
    before = {p for p in os.listdir(tmp) if p.startswith("docmat-")}
    upload(client, "/api/preview", "sample.docx", sample_docx)
    convert(client, recognize(client, "sample.docx", sample_docx), "Новый текст")
    after = {p for p in os.listdir(tmp) if p.startswith("docmat-")}
    assert after == before


@pytest.mark.parametrize("name", ["../../etc/passwd", 'a<b>:"c', ""])
def test_safe_basename(name):
    from app.main import safe_basename

    result = safe_basename(name)
    assert "/" not in result and "<" not in result and result
