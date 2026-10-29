import io

from docx import Document
from openpyxl import load_workbook

from app.services import editable


def paragraphs(data: bytes) -> list:
    return Document(io.BytesIO(data)).paragraphs


def test_docx_text_includes_paragraphs_and_tables(sample_docx):
    text = editable.text_of("docx", sample_docx)
    assert text.startswith("Техническое задание\nСервис выполняет оптическое распознавание текста")
    assert text.endswith("Параметр\tЗначение\nФормат\tPDF/A")


def test_unchanged_text_returns_same_document(sample_docx):
    result = editable.apply_text("docx", sample_docx, editable.text_of("docx", sample_docx))
    assert result.data == sample_docx and result.applied == 0


def test_edit_inside_bold_run_keeps_formatting(sample_docx):
    text = editable.text_of("docx", sample_docx).replace("оптическое", "оптическое и рукописное")
    result = editable.apply_text("docx", sample_docx, text)
    runs = paragraphs(result.data)[1].runs
    assert [(r.text, bool(r.bold)) for r in runs] == [
        ("Сервис выполняет ", False),
        ("оптическое и рукописное распознавание", True),
        (" текста и проверку орфографии документов, переданных пользователем через веб-интерфейс.", False),
    ]
    assert paragraphs(result.data)[0].runs[0].font.size.pt == 16  # остальное оформление не тронуто


def test_edit_across_runs_goes_to_first_run(sample_docx):
    text = editable.text_of("docx", sample_docx).replace("выполняет оптическое", "делает")
    result = editable.apply_text("docx", sample_docx, text)
    p = paragraphs(result.data)[1]
    assert p.text.startswith("Сервис делает распознавание текста")
    assert p.runs[1].bold and p.runs[1].text == " распознавание"


def test_edit_in_table_cell(sample_docx):
    text = editable.text_of("docx", sample_docx).replace("PDF/A", "PDF/A-2b")
    result = editable.apply_text("docx", sample_docx, text)
    assert Document(io.BytesIO(result.data)).tables[0].cell(1, 1).text == "PDF/A-2b"


def test_joining_paragraphs_is_reported_as_skipped(sample_docx):
    text = editable.text_of("docx", sample_docx).replace("задание\nСервис", "задание Сервис")
    result = editable.apply_text("docx", sample_docx, text)
    assert result.applied == 1  # пробел вставлен в конец заголовка
    assert len(paragraphs(result.data)) == 3  # абзацы не склеены: перенос строки — служебный символ


def test_text_after_edits_matches(sample_docx):
    text = editable.text_of("docx", sample_docx)
    edited = text.replace("Документы", "Файлы").replace("орфографии", "орфографии и грамматики")
    result = editable.apply_text("docx", sample_docx, edited)
    assert editable.text_of("docx", result.data) == edited


def test_xlsx_text_and_edits(sample_xlsx):
    text = editable.text_of("xlsx", sample_xlsx)
    assert text == "Наименование\tКоличество\tЦена\tСумма\nБумага\t10\t250\t2500\nКартридж\t2\t3500.5\t7001"
    edited = text.replace("Бумага", "Бумага А4").replace("\t10\t", "\t12\t").replace("7001", "1")
    result = editable.apply_text("xlsx", sample_xlsx, edited)
    ws = load_workbook(io.BytesIO(result.data)).active
    assert ws["A2"].value == "Бумага А4"
    assert ws["B2"].value == 12  # число осталось числом
    assert ws["D3"].value == "=B3*C3"  # формула не перезаписана текстом
    assert ws["A1"].font.bold  # оформление сохранено


def test_diff_is_word_level():
    edits = editable.diff("один два три", "один дважды три")
    assert edits == [(5, 8, "дважды")]


def test_edit_across_paragraph_border_is_split():
    assert editable.diff("Сокращеныя\nвремеми", "сокращения\nвремени") == [(0, 10, "сокращения"), (11, 18, "времени")]


def test_neighbour_words_in_two_paragraphs(sample_docx):
    text = editable.text_of("docx", sample_docx)
    edited = text.replace("задание\nСервис", "заданiе\nСерв1с")
    result = editable.apply_text("docx", sample_docx, edited)
    assert editable.text_of("docx", result.data) == edited
